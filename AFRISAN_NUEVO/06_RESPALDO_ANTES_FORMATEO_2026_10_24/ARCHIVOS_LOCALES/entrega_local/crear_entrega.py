from pathlib import Path
import json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parents[1]
pages=[]
def page(title,items): pages.append((title,items))
def p(t): return ('p',t)
def h(t): return ('h',t)
def table(head,rows): return ('table',(head,rows))
page('Avance del proyecto AFRISAN',[
p('Sistema Integral de Gestión y Trazabilidad de Refrigerantes'),
p('Capstone · Informe de avance · 22 de octubre de 2026'),
p('Integrante actual y responsable de continuidad: Dante Moya'),
h('Resumen ejecutivo'),
p('El proyecto cuenta con un prototipo en .NET 8 que registra refrigerantes y cilindros, consulta inventario y permite recorrer órdenes de trabajo con entrega, reemplazo físico, fotografías de pesajes, recepción y revisión. La aplicación nueva funciona con datos temporales en memoria y borradores locales en el navegador. La persistencia en Supabase, los límites semanales y mensuales y la integración con Power BI siguen pendientes.'),
p('La principal dificultad es la diferencia entre el proceso de recarga de equipos descrito inicialmente y el reemplazo físico de cilindros aclarado durante la revisión del negocio. Este cambio afecta el MER, los criterios de retorno, los cálculos y los informes. La continuidad individual exige priorizar un flujo completo y conservar evidencia verificable de cada incremento.'),
p('Este corte actualiza la documentación con la revisión del 20 de octubre. No incorpora nuevas pruebas funcionales ni acredita aceptación del cliente. El porcentaje de avance debe determinarse con criterios acordados con el docente.'),
h('Abstract'),
p('AFRISAN has a .NET 8 prototype for refrigerant inventory and physical cylinder replacement. It includes work orders, technical records, weighing photographs and administrative review. Server data remains in memory, while the browser retains local drafts. PostgreSQL integration, branch limits and Power BI are pending. The main challenge is aligning the previous equipment refill schema with the clarified replacement process. Individual continuity requires prioritizing one verified workflow and documenting evidence. This report updates the project status without claiming additional functional tests or production readiness.'),
h('Contenido'),
p('2 Contexto y proceso de negocio\n3 Avances y evidencias\n4 Dificultades y ajustes del proyecto\n5 Problemas técnicos identificados\n6 Plan de trabajo y próximos incrementos\n7 Calidad y módulo Power BI\n8 Conclusions and personal reflection\n9 Fuentes y control documental')])
page('Contexto y proceso de negocio',[
h('Necesidad del cliente'),p('AFRISAN necesita conocer qué cilindros tiene, qué refrigerante contiene cada uno, dónde se encuentran, quién los manipula y qué pesajes y autorizaciones respaldan sus movimientos. El control de kilos por sucursal busca detectar excesos y permitir su revisión por Jefatura.'),
h('Proceso aclarado durante el desarrollo'),p('La sucursal solicita reemplazar o agregar cilindros por refrigerante y kilos. El Autorizador crea y autoriza la OT y asigna un técnico. Bodega entrega los envases; el técnico registra el retirado y el instalado, sus pesos y fotografías. Bodega recibe los retirados y los entregados sin uso. El Autorizador revisa el trabajo y realiza el cierre administrativo.'),
p('Cada cilindro contiene un único tipo de refrigerante. Una OT puede involucrar varios cilindros. El instalado permanece en la sucursal; el retirado regresa a Bodega. Un agregado no requiere inventar un cilindro retirado.'),
table(['Rol','Responsabilidad actual'],[['Bodega','Registrar y consultar inventario; entregar y recibir; definir límites habilitados.'],['Técnico','Ejecutar reemplazos o agregados y registrar pesajes, fotos y observaciones.'],['Autorizador','Crear, autorizar, asignar técnico, revisar incidencias y cerrar OT.'],['Jefatura','Definir límites con prioridad y revisar antecedentes de excesos.']]),
p('Jerarquía confirmada: Bodega, Autorizador y Jefatura, de menor a mayor autoridad. Si Jefatura modifica un límite, Bodega queda bloqueada para editarlo hasta que Jefatura la habilite. Los permisos específicos del Autorizador sobre límites aún requieren definición.'),
h('Arquitectura'),p('Se conserva Blazor WebAssembly y ASP.NET Core Web API en .NET 8, con PostgreSQL/Supabase y EF Core como persistencia prevista. La nueva API usa stores en memoria. La Web utiliza IndexedDB para borradores y fotos pendientes. Esta separación permite sustituir la persistencia temporal por una integración durable.')])
page('Avances y evidencias disponibles',[
p('La solución vigente para revisión es AFRISAN_NUEVO/03_CODIGO/Afrisan.Nuevo.slnx. La Web usa el puerto 5286 y la API el 5181. Los proyectos anteriores de la raíz permanecen como antecedentes.'),
table(['Módulo','Implementación identificada','Límite vigente'],[
['Inventario','Catálogo, cilindros, tara, capacidad, pesaje inicial, ficha y QR.','Datos centrales en memoria.'],
['Órdenes','Creación manual, detalle solicitado y autorizado y un técnico asignado.','Sucursal como texto; sin importación ERP.'],
['Bodega','Entrega y recepción de envases asociadas a OT.','Custodia entre varias OT necesita corrección.'],
['Técnico','Reemplazo o agregado con dos envases diferenciados cuando corresponde.','Reglas provisionales y pruebas integrales pendientes.'],
['Fotografías','Colección por cilindro y pesaje; validaciones de tamaño y formato.','Servidor temporal; sin Storage durable.'],
['Trabajo local','Borradores, fotos y reintentos mediante IndexedDB.','Offline y concurrencia sin validación integral.'],
['Autorización','JWT, roles y restricción del técnico a sus OT.','Cuentas públicas de demostración.'],
['Revisión','Incidencias, conclusión del Autorizador y cierre administrativo.','Una conclusión no modifica un pesaje confirmado.']]),
h('Evidencia identificable'),p('TRABAJO_TECNICO.md describe el incremento del 20 de octubre. Las capturas ACCESO_PRIMERA_ENTREGA.png, RESUMEN_API_CONECTADA.png y TECNICO_ORDENES_2026_10_20.png muestran pantallas de revisión; la última corresponde a una lista vacía de OT. No prueban un reemplazo completo.'),
p('La entrega anterior documenta compilación sin errores ni advertencias. La revisión del 20 de octubre fue estática y no ejecutó pruebas funcionales. La existencia de código, de un script SQL o de una captura se mantiene diferenciada de su validación y aceptación.')])
page('Dificultades y ajustes del proyecto',[
h('Continuidad individual'),p('Dante Moya continúa como único integrante actual. El plan anterior distribuía tareas entre varios participantes y debe adaptarse a una disponibilidad informada de tres horas diarias. Las contribuciones anteriores conservan su autoría. La pauta contempla evaluación grupal y requiere acordar su aplicación individual con el docente.'),
h('Definición del negocio y rediseño de datos'),p('El diseño anterior asumía recarga de equipos y devolución del mismo cilindro despachado. El proceso aclarado contempla reemplazo físico. Por eso cambian las relaciones entre OT, envases, movimientos y pesajes, y debe revisarse la conciliación de retornos. La arquitectura aprobada sigue siendo válida como estructura, pero el nuevo modelo de datos necesita alineación.'),
h('Creación y localización de tablas'),p('Durante la preparación de Supabase aparecieron errores por esquema existente y dificultad para localizar tablas después de ejecutar scripts. También surgió un error de clave foránea al referenciar columnas sin restricción única. Se prepararon scripts y un diccionario manual de PK, FK y tipos. El SVG exportado evidencia tablas y relaciones visibles; no acredita por sí mismo triggers, políticas o conexión de la aplicación.'),
h('Inicio de la aplicación'),p('Se reportaron páginas que no cargaban, API apagada y diferencias entre la ejecución en Visual Studio y la aplicación abierta en el navegador. El repositorio contiene dos soluciones con puertos distintos. La nueva solución ofrece un perfil de inicio para Web y API, pero ambos procesos deben estar disponibles. El mensaje de un subproceso finalizado no demuestra por sí solo un error de la aplicación.'),
h('Separación y documentación'),p('Se creó AFRISAN_NUEVO para ordenar procesos, datos, código y seguimiento. La revisión intensiva quedó en una carpeta local con hallazgos verificables. Los informes con corte del 16 de octubre describían principalmente la aplicación anterior y necesitan esta actualización para coincidir con la defensa.'),
p('Ajuste de método: revisar un proceso, registrar decisiones, implementar un incremento, comprobar resultados y documentar evidencia. No se atribuyen ceremonias Scrum ni acuerdos de aceptación que no estén respaldados.')])
page('Problemas técnicos identificados',[
p('Los siguientes escenarios provienen de la lectura del código del 20 de octubre. Su reproducción y comprobación después de corregirlos forman parte del trabajo pendiente.'),
table(['Problema','Impacto y acción siguiente'],[
['Borrador sin descarte','Un borrador adicional sincronizado puede impedir terminar toda la OT. Incorporar descarte trazado.'],
['Custodia entre OT','El cierre consulta ubicación global del envase. Vincular cada entrega y recepción a su propio movimiento.'],
['Conflicto entre pestañas','Borrar una foto puede adelantar la versión local y permitir sobrescribir pesos. Exigir versión coherente en cada modificación.'],
['Entregas repetidas','Una entrega histórica puede devolver éxito sin cambiar la custodia tras una devolución. Usar identidad de movimiento.'],
['Estado de la OT','La API admite entregas después del término técnico aunque la pantalla las restringe. Alinear reglas del servidor.'],
['Caché y fecha offline','Pueden reaparecer OT antiguas; falta fecha física separada de sincronización. Conservar instancia y fechas distintas.']]),
h('Brechas para una operación real'),p('Reiniciar la API elimina inventario, OT y fotos centrales. Crear una OT puede ubicar un envase en sucursal como preparación del ejemplo. Una respuesta de rectificación registra revisión, pero no aplica un ajuste de peso. Estos comportamientos deben sustituirse antes de operar con datos reales.'),
p('El SQL no admite el nuevo rol Autorizador en su CHECK y no representa completamente el detalle solicitado/autorizado, evidencias múltiples, técnico obligatorio de OT ni límites con historial. Conectar el prototipo al SQL actual sin corregirlo trasladaría al sistema reglas incompatibles.'),
p('Los gramos autorizados se almacenan, pero todavía no generan comparación automática con los kilos reales ni con límites semanales o mensuales. Superar un límite requiere registro e investigación según el acuerdo de negocio; el bloqueo automático de ejecución no fue confirmado.')])
page('Plan de trabajo y próximos incrementos',[
p('La disponibilidad es de tres horas diarias. La secuencia siguiente expresa prioridades y condiciones de cierre; las fechas finales se reestiman con el alcance validado. La planificación de seis horas del 20 y 21 fue una propuesta previa y no se presenta como trabajo ya ejecutado.'),
table(['Orden','Incremento','Resultado necesario'],[
['1','Estabilizar el recorrido','Reemplazo completo, recepción y cierre comprobados; descarte de borradores y custodia coherentes.'],
['2','Alinear negocio y MER','OT con técnico, detalle por gas, pesajes, pareja de envases y colección de evidencias.'],
['3','Persistir el recorrido','Datos y fotos sobreviven al reinicio con acceso autenticado y transacciones consistentes.'],
['4','Usuarios reales','Identidades y roles persistidos; cuentas activas y restricciones verificadas en API.'],
['5','Límites y ajustes','Cálculo acordado, historial del límite, protección de Jefatura y revisión de excesos.'],
['6','Módulo Power BI','Informe reproducible con fuentes, medidas, actualización y acceso definidos.'],
['7','Validación restante','Conflictos, reintentos, offline y requisitos restantes del backlog según alcance vigente.']]),
h('Decisiones que desbloquean el desarrollo'),p('Definir qué kilos acumulan los límites, si se distinguen por gas y si los períodos son de calendario. Precisar la fecha del hecho que determina la semana o mes. Acordar correcciones de pesos, cierre parcial, reapertura y descarte. Definir inventario inicial en sucursal y cómo se repone un cilindro vacío.'),
h('Criterio de avance'),p('Una tarea queda terminada cuando cumple sus reglas, conserva datos y evidencia cuando corresponda, y dispone de resultados verificables. Las pantallas programadas no equivalen a historias aceptadas. El plan de ocho columnas actualizado acompaña este informe en Markdown.'),
p('Las integraciones con Defontana, certificados, semáforo y componentes de costos siguen en el alcance documental y requieren conciliación con el proceso revisado. El manual de usuario y la matriz formal de trazabilidad permanecen aplazados según la decisión de Dante.')])
page('Calidad y módulo Power BI',[
h('Validaciones prioritarias'),p('Comprobar pesos entre tara y tara más capacidad, evidencia obligatoria de ambos pesajes en reemplazos, acceso por rol, recepción antes del cierre y comportamiento de un borrador descartado. Después ampliar a varias OT, entregas repetidas, conflictos entre pestañas y reintentos offline sin duplicación.'),
p('Cada comprobación debe conservar fecha, entorno, pasos, resultado esperado y observado y evidencia. El aislamiento entre dos técnicos necesita dos identidades de prueba; la cuenta demo única no basta para demostrarlo. Las pruebas no ejecutadas se conservan como pendientes.'),
h('Datos necesarios para Power BI'),p('El acta compromete Power BI y al menos cinco KPI. Los indicadores originales incluyen kilos cargados/devueltos, costos, CO₂e, criticidad y devolución. El reemplazo físico obliga a revisar el significado de algunos de ellos. El primer informe puede mostrar inventario y OT, junto con solicitado, autorizado e instalado, manteniendo pendiente la aprobación de cualquier sustitución de KPI.'),
p('Las vistas deben conservar una fila por hecho definido y relacionarse con fecha, sucursal, refrigerante y actor. No sumar etapas distintas como consumo adicional ni multiplicar kilos al unir una operación con varias fotos. Costos requieren precio histórico aplicado y conceptos acordados; emisiones requieren masa estimada, método y GWP con fuente.'),
h('Integración y acceso'),p('El módulo todavía no está conectado. Power BI puede consultar vistas PostgreSQL con una identidad de lectura. La presentación segura dentro de la aplicación necesita resolver autenticación, permisos, licencias y HTTPS. Las tablas de usuarios y las fotografías privadas no son necesarias para los gráficos de gestión.'),
h('Riesgos de continuidad'),p('El corte de revisión encontró código nuevo y EVA2 sin seguimiento Git. Su publicación permanece pendiente por instrucción de Dante. También debe registrarse una decisión de continuidad tecnológica: .NET 8 termina soporte el 10 de noviembre de 2026, según Microsoft. La tecnología del proyecto no se cambió en esta entrega.')])
page('Conclusions and personal reflection',[
h('Individual conclusions'),p('AFRISAN has an implemented prototype for inventory and the physical replacement workflow. It separates warehouse delivery, technical work, photographic evidence, warehouse reception and administrative review. The current server stores operational data in memory, so durable persistence remains a major integration task.'),
p('The main engineering issue is the mismatch between the original equipment refill model and the clarified cylinder replacement process. Connecting the current application to the previous SQL schema would preserve incorrect assumptions. The next increment should align the data model and verify one complete workflow before adding further modules.'),
p('Power BI remains part of the project scope. Reliable reporting depends on stable branch and user identities, consistent event dates, historical prices and limits, and clearly defined quantities. The evidence currently available does not support a fixed completion percentage or a claim of production readiness.'),
h('Individual reflection'),p('Continuing this project individually requires me to make the workload visible and explain the reasons behind each technical decision. I need to keep the business meaning of each record clear, especially when distinguishing a cylinder replacement from refrigerant consumption or emissions.'),
p('My next priority is to verify the central workflow, record observed results and resolve the model inconsistencies. I also need to keep the report, backlog and code aligned so that I can defend the actual state of the work and plan the next increment within my available time.'),
p('Esta reflexión se entrega como texto propuesto para revisión personal de Dante antes de su presentación. No contiene una firma ni certifica una experiencia o resultado de prueba adicional.')])
page('Fuentes y control documental',[
h('Fuentes del proyecto'),p('1. Acta de Constitución AFRISAN corregida, carpeta Documentación del Proyecto/ACTUAL. Objetivos, requisitos, reglas y aprobación histórica.\n2. Product Backlog AFRISAN corregido, hojas Historias de Usuario y Criterios de Aceptación.\n3. Guías 2.4 y 2.6 y pauta 2.5 de Fase 2. Estructura, monitoreo y criterios académicos.\n4. AFRISAN_NUEVO/01_NEGOCIO, especialmente definición de OT, responsables, técnico, Autorizador y Power BI.\n5. AFRISAN_NUEVO/03_CODIGO/TRABAJO_TECNICO.md y código de API, Web y contratos.\n6. Revisión del 20/10: INFORME_GENERAL.md, CODIGO.md, DATOS_Y_NEGOCIO.md y CAPSTONE.md. Hallazgos con rutas, líneas, párrafos y celdas.'),
h('Referencias técnicas'),p('Microsoft. Política de soporte de .NET. https://dotnet.microsoft.com/en-us/platform/support/policy\nMicrosoft Learn. Secure embed y esquema de estrella para Power BI. https://learn.microsoft.com/en-us/power-bi/collaborate-share/service-embed-secure y https://learn.microsoft.com/en-us/power-bi/guidance/star-schema\nSupabase. Conexión PostgreSQL y buckets privados. https://supabase.com/docs/guides/database/connecting-to-postgres y https://supabase.com/docs/guides/storage/buckets/fundamentals\nConsulta técnica registrada: 20 de octubre de 2026.'),
h('Control de esta actualización'),p('Corte documental: 22 de octubre de 2026. Base de evidencia: implementación y revisión del 20/10 y documentos anteriores. La edición actualiza el relato, la planificación y el material de defensa; no modifica código, SQL ni datos en nube. Las fuentes históricas se conservan y las versiones anteriores de los documentos actualizados quedan en ANTECEDENTES.'),
p('Responsable actual: Dante Moya. Publicación pendiente por instrucción del estudiante. El tratamiento individual de la pauta y la aceptación del proceso revisado requieren constancia del docente y del cliente respectivamente.')])

# Preserve the previous versions before updating current reports.
targets=['EVIDENCIAS_GRUPALES/EVA_2/INFORME_AVANCE_FASE_2_AFRISAN.md','EVIDENCIAS_GRUPALES/EVA_2/02_Informe_Avance_Fase2_AFRISAN.docx','EVIDENCIAS_GRUPALES/EVA_2/01_PLAN_DE_TRABAJO_8_COLUMNAS_FASE_2.md']
for rel in targets:
 src=ROOT/rel; dst=OUT/'ANTECEDENTES'/src.name
 if src.exists() and not dst.exists(): shutil.copy2(src,dst)

doc=Document(); sec=doc.sections[0]; sec.page_height=Inches(11.7); sec.page_width=Inches(8.3)
sec.top_margin=sec.bottom_margin=Inches(.68); sec.left_margin=sec.right_margin=Inches(.78)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11 if name=='Normal' else 24 if name=='Title' else 16 if name=='Heading 1' else 12)
 st.paragraph_format.space_after=Pt(7); st.paragraph_format.line_spacing=1.08
doc.core_properties.author='Dante Moya';doc.core_properties.title='Avance del proyecto AFRISAN'
footer=sec.footer.paragraphs[0];footer.alignment=2
run=footer.add_run(); fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
md=[]
for i,(title,items) in enumerate(pages):
 if i: doc.add_page_break()
 doc.add_paragraph(title,'Title' if not i else 'Heading 1');md.append('# '+title)
 for kind,value in items:
  if kind=='p': doc.add_paragraph(value);md.append(value)
  elif kind=='h':doc.add_paragraph(value,'Heading 2');md.append('## '+value)
  else:
   head,rows=value;t=doc.add_table(rows=1, cols=len(head));t.autofit=False
   for c,text in zip(t.rows[0].cells,head):c.text=text
   trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
   for ri,row in enumerate(rows):
    cells=t.add_row().cells
    for c,text in zip(cells,row):c.text=text
   for ri,row in enumerate(t.rows):
    for c in row.cells:
     pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
     for edge in ['top','left','bottom','right']:
      e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
     pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17394B' if ri==0 else 'F0F4F6' if ri%2 else 'FFFFFF');pr.append(shade)
     for para in c.paragraphs:
      para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(4)
      for run in para.runs:run.font.size=Pt(11);run.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
   md+=['| '+' | '.join(head)+' |','| '+' | '.join(['---']*len(head))+' |']+['| '+' | '.join(row)+' |' for row in rows]
  md.append('')
report=OUT/'Informe_Avance_AFRISAN_2026_10_22.docx';doc.save(report)
shutil.copy2(report,ROOT/targets[1]);text='\n\n'.join(md)
(OUT/'Informe_Avance_AFRISAN_2026_10_22.md').write_text(text,encoding='utf-8');(ROOT/targets[0]).write_text(text,encoding='utf-8')

slides=[]
def slide(title,lines,note,image=None):slides.append(dict(title=title,lines=lines,note=note,image=image))
slide('Avance del proyecto AFRISAN',['Gestión y trazabilidad de refrigerantes','Dante Moya','Capstone · 22 de octubre de 2026'],'Presento el estado del proyecto y los ajustes que requiere su continuidad individual. La evidencia técnica corresponde al corte del 20 de octubre; hoy actualizo documentación y planificación.')
slide('Necesidad del cliente',['Identificar cada cilindro y su refrigerante','Conocer ubicación, peso y responsable','Respaldar entregas y reemplazos con evidencia','Revisar kilos por sucursal y autorizaciones'],'Explicar un ejemplo: una sucursal solicita cilindros, Bodega entrega y un técnico instala. La información debe relacionarse para revisar qué ocurrió. Fuente: 01_NEGOCIO/07_OT_DEFINICION_DANTE.md.')
slide('Proceso de la orden de trabajo',['1  Solicitud y autorización con técnico asignado','2  Entrega de Bodega y reemplazo en sucursal','3  Pesajes y fotos del retirado y del instalado','4  Recepción de Bodega y revisión para cierre'],'El instalado permanece en sucursal y el retirado vuelve a Bodega. Un agregado no inventa un retiro. Autorizador y Jefatura son roles diferentes. Fuente: 01_NEGOCIO/07,08,09 y10.')
slide('Continuidad individual',['Dante Moya concentra el trabajo actual','Disponibilidad informada de tres horas diarias','Planificación por incrementos verificables','Autoría histórica y arquitectura conservadas'],'La pérdida de integrantes afecta capacidad y distribución de tareas. Mantengo las contribuciones históricas y debo acordar con el docente la aplicación individual de la pauta. No atribuir mala conducta ni causas personales a quienes ya no participan.')
slide('Avance de inventario',['Catálogo de refrigerantes','Cilindros con tara, capacidad y un tipo de gas','Pesaje inicial, consulta y ficha','Código QR generado por cilindro'],'El prototipo comienza con inventario vacío y permite registrar datos de revisión. El primer peso puede quedar pendiente, decisión provisional. Los registros están en memoria. Fuente: 03_CODIGO/README.md y contratos Inventario.cs.')
slide('Avance del trabajo técnico',['OT asignadas y entrega de envases','Reemplazo o agregado con pesajes','Varias fotografías por cilindro','Recepción, incidencias y revisión'],'La captura corresponde a la lista vacía de OT del 20 de octubre. Demuestra una pantalla disponible, no un flujo completo aprobado. El código contiene funciones que aún requieren validación integral. Fuente: TRABAJO_TECNICO.md.', 'TECNICO_ORDENES_2026_10_20.png')
slide('Arquitectura y almacenamiento',['Web Blazor y API ASP.NET Core en .NET 8','Estado central temporal en memoria','Borradores y fotos locales en IndexedDB','EF Core, Supabase y Storage pendientes'],'La solución nueva usa Web5286 y API5181. La anterior usa5226/5141. Explicar la diferencia entre RAM, navegador y base central. Mostrar por qué reiniciar la API pierde datos. Fuente: Program.cs y TRABAJO_TECNICO.md.')
slide('Dificultades durante el desarrollo',['Redefinición del reemplazo físico','Errores de esquema y relaciones en Supabase','API apagada y confusión entre soluciones','Documentos con estados y responsables anteriores'],'Se resolvió la organización en AFRISAN_NUEVO, se prepararon instrucciones de inicio y diccionario PK/FK. La integración de datos todavía necesita trabajo. Fuentes: conversación y revisión DATOS_Y_NEGOCIO.md y CAPSTONE.md.')
slide('Incongruencia del modelo de datos',['El SQL anterior representa recarga de equipos','El negocio actual reemplaza cilindros físicos','Faltan relaciones y reglas del nuevo flujo','El siguiente diseño debe alinear OT y movimientos'],'Explicar que un problema de significado no se resuelve solo cambiando una FK. Faltan Autorizador en CHECK, técnico de OT, detalle, pareja de envases, fotos múltiples y límites. Fuente: DATOS_Y_NEGOCIO.md D01 a D08.')
slide('Problemas técnicos por corregir',['Borradores sin descarte pueden bloquear una OT','Custodia puede confundirse entre varias OT','Edición concurrente y caché requieren ajustes','El cálculo de límites aún no está implementado'],'Son escenarios identificados mediante inspección estática, todavía sin reproducción funcional documentada. No presentarlos como pruebas ejecutadas. Explicar prioridad C03 para la demo y luego C02/C01. Fuente: revisión CODIGO.md.')
slide('Módulo de gráficos con Power BI',['Inventario y estado de las OT','Solicitado, autorizado y realizado por sucursal','Límites, costos e indicadores ambientales','Fuentes, fórmulas y acceso por definir'],'Power BI sigue en alcance y aún no está conectado. El acta compromete cinco KPI, algunos necesitan ajuste al reemplazo físico. El gas instalado no demuestra emisión. La integración segura necesita permisos/licencias y HTTPS. Fuentes: acta, 11_MODULO_POWER_BI.md y https://learn.microsoft.com/en-us/power-bi/collaborate-share/service-embed-secure.')
slide('Próximos incrementos',['1  Comprobar y estabilizar una OT completa','2  Alinear el MER y persistir datos y fotos','3  Incorporar usuarios, límites y rectificaciones','4  Conectar Power BI y ampliar la validación'],'El orden permite trabajar solo sin desarrollar módulos desconectados. Cada incremento debe demostrar resultados. No comprometer una fecha final que no se ha estimado. Mantener requisitos restantes de Defontana y certificados sujetos al alcance vigente. Fuente: plan actualizado.')
slide('Validación y evidencia de avance',['Compilación previa documentada','Capturas y código disponibles para revisión','Pruebas integrales y aceptación pendientes','Resultados esperados y observados por registrar'],'No afirmar un porcentaje fijo de50% ni producción. Necesito evidenciar reglas, datos y resultados. Código nuevo y EVA2 siguen locales por decisión expresa, sin commits ni push. Fuentes: TRABAJO_TECNICO.md y CAPSTONE.md.')
slide('Decisiones para continuar',['Confirmar el proceso y los criterios revisados','Definir kilos, períodos y permisos de límites','Acordar correcciones y condiciones de cierre','Priorizar la persistencia y la evidencia del flujo'],'Cerrar explicando qué decisión desbloquea el trabajo: modelo alineado y un recorrido durable. Puedo responder sobre roles, API, RAM y Power BI. La continuidad de .NET8 debe revisarse por soporte10/11/2026: https://dotnet.microsoft.com/en-us/platform/support/policy. Arquitectura actual conservada.')
(OUT/'.local/slides.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
guion=['# Guion para la defensa de AFRISAN','Dante Moya · 22 de octubre de 2026','Duración sugerida: 10 a 12 minutos, ajustable al tiempo asignado.']
for i,s in enumerate(slides):guion+=['\n## '+str(i+1)+' '+s['title'],s['note']]
(OUT/'Guion_Defensa_AFRISAN.md').write_text('\n\n'.join(guion),encoding='utf-8')
plan='# Plan de trabajo actualizado al 22 de octubre de 2026\n\nResponsable actual: Dante Moya. Capacidad: tres horas diarias. Estimaciones por sesión de trabajo, sujetas a revisión. Los resultados no ejecutados permanecen pendientes.\n\n| Competencia o área | Actividad | Recursos | Tiempo | Responsable | Avance observado | Estado | Ajuste o siguiente acción |\n|---|---|---|---|---|---|---|---|\n'
rows=[['Gestión','Actualizar informe y defensa','Informe, PPT y revisión','Edición del 22/10','Dante Moya','Documentación actualizada al nuevo prototipo','En curso','Revisar personalmente reflexión y ensayar'],['Software','Estabilizar OT y custodia','API y Blazor','Sesiones de 3 h; reestimar','Dante Moya','Recorrido programado con defectos estáticos identificados','En curso','Reproducir C03 y corregir descarte'],['Datos','Alinear MER y reglas','SQL y procesos','Primera sesión de diseño de 3 h','Dante Moya','SQL anterior incompatible con reemplazo','En curso','Definir identidades, pesajes y fotos'],['Datos','Persistir flujo y evidencias','EF Core, Supabase, Storage','Por estimar después del modelo','Dante Moya','Nueva aplicación en memoria','No iniciado','Validar conservación tras reinicio'],['Seguridad','Usuarios y roles reales','API e identidad persistida','Por estimar','Dante Moya','JWT demo con restricciones','En curso','Definir una estrategia de autenticación'],['Negocio','Límites y rectificaciones','Reglas y API','Por estimar tras acuerdos','Dante Moya','Reglas documentadas; fórmula pendiente','No iniciado','Definir kilos, período y protección'],['Analítica','Módulo Power BI','Vistas y Power BI','Por estimar tras datos','Dante Moya','Requisito y propuesta documentados','No iniciado','Acordar medidas y forma de acceso'],['Calidad','Validar flujo y permisos','Casos y evidencias','Sesiones de 3 h','Dante Moya','Sin prueba integral documentada','No iniciado','Registrar esperado y observado']]
plan+='\n'.join('| '+' | '.join(r)+' |' for r in rows)
(ROOT/targets[2]).write_text(plan,encoding='utf-8');(OUT/'Plan_Trabajo_Actualizado.md').write_text(plan,encoding='utf-8')
print('Word, fuentes Markdown, guion y plan creados. Diapositivas:',len(slides))
